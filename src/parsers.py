"""Document text extractors for PDF, DOCX, TXT, MD."""
import os
from typing import Tuple


def extract_text(filepath: str) -> Tuple[str, int]:
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        return _extract_pdf(filepath)
    if ext == ".docx":
        return _extract_docx(filepath)
    if ext in (".txt", ".md", ".markdown"):
        return _extract_txt(filepath)
    raise ValueError(f"Unsupported file type: {ext} (supported: .pdf, .docx, .txt, .md)")


def _extract_txt(filepath: str) -> Tuple[str, int]:
    with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()
    return text.strip(), len(text)


def _extract_pdf(filepath: str) -> Tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(filepath)
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    text = "\n".join(parts).strip()
    return text, len(text)


def _extract_docx(filepath: str) -> Tuple[str, int]:
    from docx import Document

    doc = Document(filepath)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    text = "\n".join([p for p in parts if p.strip()]).strip()
    return text, len(text)
